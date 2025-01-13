import datetime
import os
import openpyxl

print ("orientações:\n -digitar nome completo\n -toda vez que for bater o ponto, digitar o nome da mesma forma da primeira vez\n -duvidas, entrar em contato no ramal ****")



usuario = int(input(f"opções:\n [1]novo usuário \n [2]usuário já existente\n Opção:".upper()))
if usuario == 1:
    nome = input("nome completo: ".upper())
    nome_arquivo = f"\\\\kali\\suporte\\ponto_{nome}.xlsx"
    # workbook = None
    if not os.path.exists(nome_arquivo):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Folha de ponto"
        sheet.append(["Ação", "Data e Hora "])

    else:
        workbook = openpyxl.load_workbook(nome_arquivo)
    


elif usuario == 2:
    nomes_diretorio = os.listdir("\\\\kali\\suporte")
    arquivos = [f for f in nomes_diretorio if f.endswith (".xlsx")]
    if arquivos:
        print ("Usuarios existentes")
        dicionario_user = {}
    for i, file in enumerate (nomes_diretorio):
        dicionario_user[i + 1]= file
        print (f"[{i+1}]{file}")
    
        selecionar_usuario = int(input("Selecione o numero que corresponde ao seu usuario:\n"))
        if selecionar_usuario in dicionario_user:
            usuario_selecionado = dicionario_user[selecionar_usuario]
            nome_arquivo = f"\\\\kali\\suporte\\{usuario_selecionado}"
            workbook = openpyxl.load_workbook(nome_arquivo)
        else:
            print ("Número de usuário inválido!")
    else: 
        print ("Nenhum arquivo de usuario existente foi encontrado!!")
else:
    print("opção inválida!")
    exit()

def ver_ultima_acao(workbook):
    sheet = workbook["Folha de ponto"]
    ultima_acao_cell = sheet.cell(row = sheet.max_row, column = 1)
    return ultima_acao_cell.value

def ponto(acao, workbook):
    now = datetime.datetime.now()
    data_hora = now.strftime("%Y/%m/%d -- %H:%M:%S")

    # workbook = openpyxl.load_workbook(nome_arquivo)
    sheet = workbook["Folha de ponto"]
    sheet.append([acao, data_hora])

    workbook.save(nome_arquivo)
    print(f"{acao}: {data_hora}")


ultima_acao = ver_ultima_acao(workbook) if workbook else None

while True:
    if ultima_acao == "Entrada":
        acao = "Saida"
    else:
        acao = "Entrada"

    input("Pressione a tecla Enter para registrar o ponto!!")
    ponto(acao, workbook)
    ultima_acao = acao